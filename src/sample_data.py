import os
from pathlib import Path
from config import SAMPLE_DOCS_DIR

def generate_sample_documents(target_dir: Path = SAMPLE_DOCS_DIR):
    """Generates a complete suite of realistic approved college documents across all 5 file formats."""
    target_dir.mkdir(parents=True, exist_ok=True)
    created_files = []

    # 1. Generate Academic Regulations Handbook (PDF)
    pdf_path = target_dir / "Academic_Regulations_Handbook_2024.pdf"
    if not pdf_path.exists():
        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib import colors

        doc = SimpleDocTemplate(str(pdf_path), pagesize=letter)
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'DocTitle',
            parent=styles['Heading1'],
            fontSize=20,
            leading=24,
            textColor=colors.HexColor('#1E3A8A'),
            spaceAfter=14
        )
        heading_style = ParagraphStyle(
            'DocHeading',
            parent=styles['Heading2'],
            fontSize=14,
            leading=18,
            textColor=colors.HexColor('#1D4ED8'),
            spaceAfter=8
        )
        body_style = ParagraphStyle(
            'DocBody',
            parent=styles['Normal'],
            fontSize=10,
            leading=14,
            spaceAfter=6
        )

        story = []

        # --- Page 1: Grading System & Academic Standards ---
        story.append(Paragraph("University Academic Regulations Handbook 2024-2025", title_style))
        story.append(Paragraph("Chapter 1: Academic Grading System and Standards", heading_style))
        story.append(Paragraph(
            "1.1 Grading Scale: The university follows a 10-point Cumulative Grade Point Average (CGPA) system. "
            "Grades are assigned as follows: Grade O (Outstanding: 90-100%, 10 points), Grade A+ (Excellent: 80-89%, 9 points), "
            "Grade A (Very Good: 70-79%, 8 points), Grade B+ (Good: 60-69%, 7 points), Grade B (Above Average: 55-59%, 6 points), "
            "Grade C (Average: 50-54%, 5 points), Grade P (Pass: 40-49%, 4 points), and Grade F (Fail: Below 40%, 0 points).",
            body_style
        ))
        story.append(Paragraph(
            "1.2 Graduation Requirements: An undergraduate student must complete a minimum of 160 academic credits "
            "with a cumulative GPA of not less than 5.0 to be eligible for the award of the Bachelor degree.",
            body_style
        ))
        story.append(Spacer(1, 20))
        story.append(PageBreak())

        # --- Page 2: Attendance Policies & Condonation ---
        story.append(Paragraph("Chapter 2: Attendance Policies and Condonation Rules", heading_style))
        story.append(Paragraph(
            "2.1 Mandatory Minimum Attendance: Every student must maintain a minimum attendance of 75% in each registered course "
            "(theory and laboratory separately) to be eligible to appear for the End-Semester Examination.",
            body_style
        ))
        story.append(Paragraph(
            "2.2 Medical Condonation: In cases of prolonged genuine medical illness or hospitalization, the Dean of Academic Affairs "
            "may grant condonation of attendance shortage up to a maximum of 10% (i.e. between 65% and 74% attendance). "
            "A condonation fee of ₹500 per subject must be remitted along with a valid Government Medical Certificate within 5 working days of recovery.",
            body_style
        ))
        story.append(Paragraph(
            "2.3 Severe Shortage (Below 65%): Students whose attendance falls below 65% under any circumstances are strictly ineligible "
            "for the end-semester examinations and shall be awarded a 'Detained' (W) grade. They must re-register for the course in the subsequent semester.",
            body_style
        ))
        story.append(Spacer(1, 20))
        story.append(PageBreak())

        # --- Page 3: Examinations, Re-evaluation & Malpractice ---
        story.append(Paragraph("Chapter 3: Examination Regulations, Re-evaluation & Code of Conduct", heading_style))
        story.append(Paragraph(
            "3.1 Re-evaluation Procedure: A student dissatisfied with the semester examination evaluation may apply for re-evaluation "
            "within 14 calendar days of results announcement. The non-refundable re-evaluation fee is ₹1,000 per theory course. "
            "If the re-evaluation results in a grade change of two or more levels, 50% of the fee will be refunded.",
            body_style
        ))
        story.append(Paragraph(
            "3.2 Academic Malpractice: Any student found possessing prohibited electronic gadgets, notes, or unauthorized aids inside "
            "the examination hall will have their paper cancelled immediately and shall face suspension from university examinations for one academic year.",
            body_style
        ))
        story.append(Paragraph(
            "3.3 Make-Up Examination: A make-up examination will only be conducted for students who missed the regular examination due to "
            "official university representation in national sports or documented severe medical hospitalization approved by the Controller of Examinations.",
            body_style
        ))

        doc.build(story)
        created_files.append(pdf_path.name)

    # 2. Generate Examination Notice (TXT)
    txt_path = target_dir / "Examination_Notice_Fall2024.txt"
    if not txt_path.exists():
        notice_content = """===============================================================
OFFICE OF THE CONTROLLER OF EXAMINATIONS
CIRCULAR / OFFICIAL NOTICE: FALL 2024 SEMESTER EXAMINATIONS
Notice Reference Number: COE/NOT/2024/FALL/089
Date of Release: October 15, 2024
Target Audience: All Registered Students and Academic Faculty
===============================================================

1. MID-TERM EXAMINATION SCHEDULE:
The Fall 2024 Mid-Semester Examinations for all Undergraduate and Postgraduate programs will commence on Monday, November 18, 2024, and conclude on Monday, November 25, 2024.
Examinations will be held in two sessions daily:
- Morning Session: 09:30 AM to 12:30 PM
- Afternoon Session: 02:00 PM to 05:00 PM

2. MANDATORY FEE CLEARANCE DEADLINE:
All semester tuition and examination fees must be cleared in full on or before Tuesday, November 5, 2024, 05:00 PM IST.
Students with uncleared dues after November 5, 2024 will not be issued Examination Hall Tickets under any circumstances.

3. HALL TICKET ISSUANCE:
Digital Hall Tickets will be made available for download on the Student Portal starting Sunday, November 10, 2024.
Students must bring a printed physical copy of their Hall Ticket and valid University Student ID Card for every examination session.

4. PROHIBITED ITEMS IN EXAMINATION HALLS:
Strict vigilance will be enforced by flying squads. The following items are strictly prohibited:
- Mobile phones, smartwatches, fitness bands, and Bluetooth devices.
- Programmable calculators (only non-programmable scientific calculators like fx-991MS are permitted for engineering papers).
- Bags, books, blank papers, and smart glasses.
Possession of any prohibited device will be treated as Level-2 Malpractice.

5. CONTACT & GRIEVANCE HELPDESK:
For hall ticket discrepancies or exam timetable clashes, email: coe-helpdesk@university.edu or visit Room 204, Administrative Block (Office Hours: 10:00 AM - 04:00 PM).
"""
        with open(txt_path, "w", encoding="utf-8") as f:
            f.write(notice_content)
        created_files.append(txt_path.name)

    # 3. Generate Course Syllabus (DOCX)
    docx_path = target_dir / "CS101_Data_Structures_Syllabus.docx"
    if not docx_path.exists():
        import docx
        doc = docx.Document()
        doc.add_heading("Department of Computer Science & Engineering", level=1)
        doc.add_heading("Course Syllabus: CS101 - Data Structures and Algorithms", level=2)
        doc.add_paragraph("Academic Year: 2024-2025 | Course Credits: 4 (3 Lectures + 2 Lab Hours/Week)")
        doc.add_paragraph("Course Instructor: Dr. Rajesh Kumar (Email: rkumar@university.edu | Office: CS-302)")

        doc.add_heading("1. Course Description & Learning Objectives", level=3)
        doc.add_paragraph(
            "This course covers fundamental data structures, algorithm design principles, and computational complexity. "
            "Students will master arrays, linked lists, stacks, queues, binary search trees, AVL trees, heaps, hashing techniques, "
            "and graph algorithms including Dijkstra's algorithm and Prim's algorithm."
        )

        doc.add_heading("2. Recommended Textbooks and References", level=3)
        doc.add_paragraph(
            "• Primary Textbook: 'Introduction to Algorithms' by Thomas H. Cormen, Charles E. Leiserson, Ronald L. Rivest, and Clifford Stein (MIT Press, 4th Edition).\n"
            "• Reference 1: 'Data Structures and Algorithm Analysis in C++' by Mark Allen Weiss (Pearson, 4th Edition).\n"
            "• Reference 2: 'Algorithms' by Robert Sedgewick and Kevin Wayne (Addison-Wesley Professional)."
        )

        doc.add_heading("3. Evaluation Weightage and Grading Policy", level=3)
        table = doc.add_table(rows=1, cols=3)
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "Assessment Component"
        hdr_cells[1].text = "Weightage (%)"
        hdr_cells[2].text = "Schedule"

        components = [
            ("Continuous Internal Assessments (2 Quizzes)", "15%", "Week 4 and Week 9"),
            ("Mid-Semester Written Examination", "25%", "Week 8"),
            ("Weekly Hands-on Programming Labs", "20%", "Continuous evaluation"),
            ("Final Capstone Algorithm Project", "10%", "Submission Week 14"),
            ("End-Semester Final Examination", "30%", "Week 16")
        ]
        for comp, wt, sch in components:
            row_cells = table.add_row().cells
            row_cells[0].text = comp
            row_cells[1].text = wt
            row_cells[2].text = sch

        doc.add_heading("4. Late Submission Policy", level=3)
        doc.add_paragraph(
            "Programming assignments submitted after the deadline will incur a 10% penalty per day up to a maximum of 3 days. "
            "No assignments will be accepted 72 hours beyond the posted deadline."
        )

        doc.save(str(docx_path))
        created_files.append(docx_path.name)

    # 4. Generate Attendance & Grading Matrix (Excel)
    xlsx_path = target_dir / "Attendance_and_Grading_Matrix.xlsx"
    if not xlsx_path.exists():
        import openpyxl
        wb = openpyxl.Workbook()

        # Sheet 1: Grading Scale
        ws_grades = wb.active
        ws_grades.title = "Grading_Scale"
        ws_grades.append(["Grade", "Min_Score", "Max_Score", "Grade_Points", "Performance_Remark", "Academic_Status"])
        grading_data = [
            ["O", 90, 100, 10, "Outstanding Performance", "Distinction"],
            ["A+", 80, 89, 9, "Excellent Performance", "First Class with Distinction"],
            ["A", 70, 79, 8, "Very Good Performance", "First Class"],
            ["B+", 60, 69, 7, "Good Performance", "First Class"],
            ["B", 55, 59, 6, "Above Average Performance", "Second Class"],
            ["C", 50, 54, 5, "Average Performance", "Second Class"],
            ["P", 40, 49, 4, "Passing Standard", "Pass"],
            ["F", 0, 39, 0, "Failing Standard", "Repeat / Remedial Required"]
        ]
        for row in grading_data:
            ws_grades.append(row)

        # Sheet 2: Attendance Rules & Fines
        ws_att = wb.create_sheet(title="Attendance_Rules")
        ws_att.append(["Attendance_Range", "Status", "Eligibility_for_Final_Exam", "Fine_Amount_INR", "Required_Documentation"])
        attendance_data = [
            ["85% - 100%", "Exemplary", "Fully Eligible with Merit Certificate", 0, "None"],
            ["75% - 84%", "Satisfactory", "Eligible to write Exam", 0, "None"],
            ["65% - 74%", "Condonation Bracket", "Eligible only upon Dean Approval & Condonation", 500, "Official Hospital/Doctor Medical Certificate"],
            ["Below 65%", "Detained (W Grade)", "Strictly Ineligible - Course Repeat Required", 0, "Academic Counseling Notice"]
        ]
        for row in attendance_data:
            ws_att.append(row)

        wb.save(str(xlsx_path))
        created_files.append(xlsx_path.name)

    # 5. Generate Campus Facilities & Hostel Guide (PPTX)
    pptx_path = target_dir / "Campus_Facilities_and_Hostel_Guide.pptx"
    if not pptx_path.exists():
        from pptx import Presentation
        from pptx.util import Inches, Pt

        prs = Presentation()

        # Slide 1: Title Slide
        slide_layout_title = prs.slide_layouts[0]
        s1 = prs.slides.add_slide(slide_layout_title)
        s1.shapes.title.text = "Campus Facilities & Student Services Guide 2024"
        s1.placeholders[1].text = "Approved Guidelines for Students, Faculty & Hostel Residents\nOffice of Student Affairs"

        # Slide 2: Central Library Hours
        slide_layout_bullet = prs.slide_layouts[1]
        s2 = prs.slides.add_slide(slide_layout_bullet)
        s2.shapes.title.text = "Central Library Timings & Access Rules"
        tf2 = s2.placeholders[1].text_frame
        tf2.text = "Library Operating Hours:"
        p2_1 = tf2.add_paragraph()
        p2_1.text = "• Monday through Friday: 08:00 AM to 10:00 PM (Reading halls open until 11:30 PM during exam weeks)."
        p2_2 = tf2.add_paragraph()
        p2_2.text = "• Weekends (Saturday & Sunday): 09:00 AM to 06:00 PM."
        p2_3 = tf2.add_paragraph()
        p2_3.text = "• Borrowing Privileges: Undergraduates may borrow up to 4 books for 14 days. Renewal allowed once online."
        p2_4 = tf2.add_paragraph()
        p2_4.text = "• Overdue Penalty: ₹5 per book per day of delay."

        # Slide 3: Hostel Curfew & Mess Timings
        s3 = prs.slides.add_slide(slide_layout_bullet)
        s3.shapes.title.text = "Hostel Rules, Curfew & Dining Services"
        tf3 = s3.placeholders[1].text_frame
        tf3.text = "Residential Standards for Hostels A, B, C and D:"
        p3_1 = tf3.add_paragraph()
        p3_1.text = "• Campus In-Time / Curfew: 09:30 PM sharp on all days. Late entry requires written permission from Chief Warden."
        p3_2 = tf3.add_paragraph()
        p3_2.text = "• Dining Hall Hours: Breakfast 07:30-09:00 AM | Lunch 12:30-02:00 PM | Dinner 07:30-09:30 PM."
        p3_3 = tf3.add_paragraph()
        p3_3.text = "• Night Leave Pass: Must be submitted on Student ERP at least 24 hours in advance with parent SMS confirmation."

        # Slide 4: Health Center & Emergency Contacts
        s4 = prs.slides.add_slide(slide_layout_bullet)
        s4.shapes.title.text = "Health Center & Emergency Helplines"
        tf4 = s4.placeholders[1].text_frame
        tf4.text = "24/7 Emergency Support Services:"
        p4_1 = tf4.add_paragraph()
        p4_1.text = "• Campus Medical Health Center: Open 24/7 with resident medical officer and pharmacy in Block E."
        p4_2 = tf4.add_paragraph()
        p4_2.text = "• Emergency Ambulance Direct Line: Campus Extension 4444 or Mobile +91-98765-43210."
        p4_3 = tf4.add_paragraph()
        p4_3.text = "• Psychological Counseling & Mental Wellness Helpdesk: Ext 4110 (Confidential consultations Mon-Fri 10am-5pm)."
        p4_4 = tf4.add_paragraph()
        p4_4.text = "• Campus Security Control Room: Ext 1000 / 1001 (Available 24 hours)."

        prs.save(str(pptx_path))
        created_files.append(pptx_path.name)

    return created_files
